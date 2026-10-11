#!/usr/bin/env python3
"""Export a small CHS-DRG 3.0 RulePack slice for Vitest golden cases.

Reads src-tauri/resources/rules/chs-drg-3.0.xlsx and writes
src/lib/engine/fixtures/chs-drg-3.0-slice.json

Only keeps MDCA/MDCF, ADRG FM3/FR1, DRG FM39/FR19, related set membership
for I21.001 / 36.0601, and one real CC (+ minimal exclude slice).
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover
    raise SystemExit("需要 openpyxl：pip install openpyxl") from exc

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "src-tauri" / "resources" / "rules" / "chs-drg-3.0.xlsx"
OUT = ROOT / "src" / "lib" / "engine" / "fixtures" / "chs-drg-3.0-slice.json"

KEEP_MDC = {"MDCA", "MDCF"}
KEEP_ADRG = {"FM3", "FR1"}
KEEP_DRG = {"FM39", "FR19"}
KEEP_SET_CODES = {
    "DI_F00": {"I21.001", "I21.002"},
    "OP_FM3": None,  # keep all official members (small)
    "DI_FR1": {"I21.001", "I21.002", "I21.003"},
}
KEEP_CC = {"E11.101"}
# Only keep principal-relevant / self codes from the CC exclude set (not full EX_23).
KEEP_EXCLUDE_CODES = {"EX_23": {"E11.101", "E11.100X051", "E10.000"}}


def find_header(ws, required: list[str]):
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        vals = [str(c).strip() if c is not None else "" for c in row]
        idx = {v: j for j, v in enumerate(vals) if v}
        if all(h in idx for h in required):
            return i, idx
    raise ValueError(f"sheet `{ws.title}` missing headers: {required}")


def cell(row, idx: dict, name: str) -> str:
    j = idx.get(name)
    if j is None:
        return ""
    v = row[j]
    if v is None:
        return ""
    return str(v).strip()


def skip_code(code: str) -> bool:
    t = code.strip()
    return not t or t in {"000", "0000"}


def main() -> int:
    if not XLSX.exists():
        print(f"missing xlsx: {XLSX}", file=sys.stderr)
        return 1

    wb = load_workbook(XLSX, read_only=True, data_only=True)

    # --- MDC ---
    ws = wb["MDC"]
    hr, idx = find_header(ws, ["MDC编码"])
    mdcs = []
    for row in ws.iter_rows(min_row=hr + 1, values_only=True):
        code = cell(row, idx, "MDC编码")
        if code not in KEEP_MDC or skip_code(code):
            continue
        mdcs.append(
            {
                "code": code,
                "name": cell(row, idx, "MDC名称"),
                "ruleExpr": cell(row, idx, "MDC规则"),
                "sort": int(float(cell(row, idx, "排序") or "0")),
            }
        )

    # --- ADRG ---
    ws = wb["ADRG"]
    hr, idx = find_header(ws, ["ADRG编码"])
    adrgs = []
    for row in ws.iter_rows(min_row=hr + 1, values_only=True):
        code = cell(row, idx, "ADRG编码")
        if code not in KEEP_ADRG or skip_code(code):
            continue
        adrgs.append(
            {
                "code": code,
                "name": cell(row, idx, "ADRG名称"),
                "ruleExpr": cell(row, idx, "ADRG规则"),
                "mdcCode": cell(row, idx, "所属MDC编码"),
                "sort": int(float(cell(row, idx, "MDC内排序") or "0")),
            }
        )

    # --- DRG ---
    ws = wb["DRG"]
    hr, idx = find_header(ws, ["DRG"])
    drgs = []
    for row in ws.iter_rows(min_row=hr + 1, values_only=True):
        code = cell(row, idx, "DRG")
        if code not in KEEP_DRG or skip_code(code):
            continue
        drgs.append(
            {
                "code": code,
                "name": cell(row, idx, "DRG名称"),
                "ruleExpr": cell(row, idx, "DRG规则"),
                "adrgCode": cell(row, idx, "所属ADRG编码"),
                "mdcCode": cell(row, idx, "所属MDC编码"),
                "sort": int(float(cell(row, idx, "排序") or "0")),
            }
        )

    # --- code sets ---
    ws = wb["集合"]
    hr, idx = find_header(ws, ["集合编号", "ICD编码"])
    sets: dict[str, list[dict[str, str]]] = {sid: [] for sid in KEEP_SET_CODES}
    for row in ws.iter_rows(min_row=hr + 1, values_only=True):
        sid = cell(row, idx, "集合编号")
        if sid not in KEEP_SET_CODES:
            continue
        icd = cell(row, idx, "ICD编码").upper()
        if skip_code(icd):
            continue
        allow = KEEP_SET_CODES[sid]
        if allow is not None and icd not in allow:
            continue
        sets[sid].append(
            {
                "icdCode": icd,
                "icdName": cell(row, idx, "ICD名称"),
                "type": cell(row, idx, "类型编号"),
            }
        )

    # --- CC ---
    ws = wb["CC"]
    hr, idx = find_header(ws, ["疾病编码"])
    cc = []
    for row in ws.iter_rows(min_row=hr + 1, values_only=True):
        icd = cell(row, idx, "疾病编码").upper()
        if icd not in KEEP_CC:
            continue
        kind = cell(row, idx, "类型").upper()
        if kind not in {"CC", "MCC"}:
            continue
        cc.append(
            {
                "code": icd,
                "name": cell(row, idx, "疾病名称"),
                "kind": kind,
                "excludeSet": cell(row, idx, "排除表"),
            }
        )

    # --- exclude sets (minimal) ---
    ws = wb["排除表"]
    hr, idx = find_header(ws, ["集合编号", "ICD编码"])
    exclude_sets: dict[str, list[dict[str, str]]] = {
        sid: [] for sid in KEEP_EXCLUDE_CODES
    }
    for row in ws.iter_rows(min_row=hr + 1, values_only=True):
        sid = cell(row, idx, "集合编号")
        if sid not in KEEP_EXCLUDE_CODES:
            continue
        icd = cell(row, idx, "ICD编码").upper()
        if skip_code(icd):
            continue
        if icd not in KEEP_EXCLUDE_CODES[sid]:
            continue
        exclude_sets[sid].append(
            {"icdCode": icd, "icdName": cell(row, idx, "ICD名称")}
        )

    wb.close()

    # sanity
    assert any(m["code"] == "MDCF" for m in mdcs), "MDCF missing"
    assert any(a["code"] == "FM3" for a in adrgs), "FM3 missing"
    assert any(d["code"] == "FM39" for d in drgs), "FM39 missing"
    assert any(x["icdCode"] == "I21.001" for x in sets["DI_F00"]), "I21.001 not in DI_F00"
    assert any(x["icdCode"] == "36.0601" for x in sets["OP_FM3"]), "36.0601 not in OP_FM3"
    assert any(x["icdCode"] == "I21.001" for x in sets["DI_FR1"]), "I21.001 not in DI_FR1"
    assert any(c["code"] == "E11.101" for c in cc), "CC E11.101 missing"

    payload = {
        "versionId": "chs-drg-3.0-slice",
        "name": "CHS-DRG 3.0（金标切片）",
        "mdcs": sorted(mdcs, key=lambda x: x["sort"]),
        "adrgs": sorted(adrgs, key=lambda x: (x["mdcCode"], x["sort"])),
        "drgs": sorted(drgs, key=lambda x: (x["adrgCode"], x["sort"])),
        "sets": sets,
        "cc": cc,
        "excludeSets": exclude_sets,
        "meta": {
            "source": "src-tauri/resources/rules/chs-drg-3.0.xlsx",
            "note": "Subset for Vitest golden cases; not the full rule pack.",
        },
    }

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    set_members = sum(len(v) for v in sets.values())
    print(
        f"wrote {OUT.relative_to(ROOT)} "
        f"(mdc={len(mdcs)} adrg={len(adrgs)} drg={len(drgs)} "
        f"set_members={set_members} cc={len(cc)})"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
